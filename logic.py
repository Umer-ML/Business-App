"""
logic.py - pure business logic (NO Streamlit, NO database in here).

Everything in this file works only on pandas DataFrames, so it is easy to
test:  python -m pytest tests/
Contains: constants, calculations (balance/status), Excel + PDF export,
and the Excel/CSV/PDF import engine.
"""

from __future__ import annotations

import io
import re
from datetime import date, datetime
from xml.sax.saxutils import escape

import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill

# =====================================================================
# CONSTANTS
# =====================================================================
BUSINESS_NAME = "Business Management System"  # printed on every PDF - change to your company name
CUR = "Rs."
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

COLUMNS = [
    "Txn ID", "Date", "Bill No", "Party / Company", "Product / Item", "Qty",
    "Subtotal", "Discount", "Tax", "Bill Amount", "Paid Amount", "Return Amount",
    "Balance", "Payment Method", "Cheque No", "Cheque Date", "Bank Name",
    "Cheque Status", "Sale Base", "Status", "Remarks",
]
ITEM_COLS = ["Txn ID", "Product", "Qty", "Rate", "Amount"]
PRODUCT_COLS = ["Product", "Unit", "Rate"]
NUM_COLS = ["Qty", "Subtotal", "Discount", "Tax", "Bill Amount", "Paid Amount", "Return Amount"]
DATE_COLS = ["Date", "Cheque Date"]
TEXT_COLS = ["Txn ID", "Bill No", "Party / Company", "Product / Item", "Payment Method", "Cheque No",
             "Bank Name", "Cheque Status", "Sale Base", "Remarks"]
MONEY = {"Subtotal", "Discount", "Tax", "Bill Amount", "Paid Amount", "Return Amount", "Balance", "Rate", "Amount",
         "Running Balance", "Total Outstanding", "Current (0-30d)", "31-60 Days", "61-90 Days", "90+ Days"}

PAYMENT_METHODS = ["Not specified", "Cash", "Cheque", "Bank Transfer", "Other"]
CHEQUE_STATUSES = ["", "Pending", "Cleared", "Bounced"]
STATUS_OPTIONS = ["UNPAID", "PARTIALLY PAID", "PAID", "CHEQUE PENDING", "RETURNED"]
NEW_PARTY = "+ Add a new party..."
NO_PRODUCT = "Select product"

GROUPS = [
    ("Overview", ["Dashboard"]),
    ("Transactions", ["New Transaction", "Bills & Payments", "Cheque Tracker"]),
    ("Records", ["Party Ledger", "Parties Directory", "Search Center", "Products", "Master Database"]),
    ("Intelligence", ["Reports", "Activity Log"]),
    ("System", ["Import Data", "User Accounts"]),
]


# =====================================================================
# CALCULATIONS
# =====================================================================
def _txt(s: pd.Series) -> pd.Series:
    out = s.where(s.notna(), "").astype(str).str.strip().replace("nan", "")
    # remove invisible control characters (they crash Excel export, common in PDF imports)
    return out.str.replace(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", regex=True)


def _num(s: pd.Series) -> pd.Series:
    cleaned = s.astype(str).str.replace(r"[^\d.\-]", "", regex=True)
    return pd.to_numeric(cleaned, errors="coerce").fillna(0.0)


def _date(s: pd.Series) -> pd.Series:
    """Parse dates row by row: 2026-01-05 (ISO) is read as-is, 05/01/2026 is read day-first."""
    if pd.api.types.is_datetime64_any_dtype(s):
        return s
    txt = s.astype(str).str.strip()
    iso = pd.to_datetime(txt.str[:10], format="%Y-%m-%d", errors="coerce")
    other = pd.to_datetime(txt, errors="coerce", dayfirst=True, format="mixed")
    return iso.fillna(other)


def recompute(df: pd.DataFrame) -> pd.DataFrame:
    """Normalise types and recalculate Balance + Status for every row."""
    df = df.copy()
    for c in COLUMNS:
        if c not in df.columns:
            df[c] = np.nan
    for c in NUM_COLS:
        df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0.0).astype(float)
    for c in DATE_COLS:
        df[c] = _date(df[c])
    for c in TEXT_COLS:
        df[c] = _txt(df[c])
    df.loc[df["Subtotal"] == 0, "Subtotal"] = df["Bill Amount"]
    df["Balance"] = (df["Bill Amount"] - df["Paid Amount"] - df["Return Amount"]).round(2)
    b, r, bal, paid = df["Bill Amount"], df["Return Amount"], df["Balance"], df["Paid Amount"]
    df["Status"] = np.select(
        [(b > 0) & (r >= b), df["Cheque Status"].eq("Pending"), bal <= 0, paid > 0],
        ["RETURNED", "CHEQUE PENDING", "PAID", "PARTIALLY PAID"], default="UNPAID",
    )
    return df[COLUMNS].reset_index(drop=True)


def _empty(cols):
    return pd.DataFrame(columns=cols)


def compute_aging(df: pd.DataFrame) -> pd.DataFrame:
    """Calculates aged receivables: Current (0-30 days), 31-60 days, 61-90 days, 90+ days."""
    cols = ["Party / Company", "Current (0-30d)", "31-60 Days", "61-90 Days", "90+ Days", "Total Outstanding"]
    if df.empty:
        return pd.DataFrame(columns=cols)
    out = df[df["Balance"] > 0].copy()
    if out.empty:
        return pd.DataFrame(columns=cols)
    today = pd.Timestamp(date.today())
    out["Days"] = (today - out["Date"]).dt.days.fillna(0).astype(int)
    out["Current (0-30d)"] = np.where(out["Days"] <= 30, out["Balance"], 0.0)
    out["31-60 Days"] = np.where(out["Days"].between(31, 60), out["Balance"], 0.0)
    out["61-90 Days"] = np.where(out["Days"].between(61, 90), out["Balance"], 0.0)
    out["90+ Days"] = np.where(out["Days"] > 90, out["Balance"], 0.0)
    
    aging = out.groupby("Party / Company")[["Current (0-30d)", "31-60 Days", "61-90 Days", "90+ Days", "Balance"]].sum().reset_index()
    aging = aging.rename(columns={"Balance": "Total Outstanding"})
    return aging.sort_values("Total Outstanding", ascending=False).reset_index(drop=True)


def compute_party_statement(df: pd.DataFrame, party: str | None = None) -> pd.DataFrame:
    """Computes a proper ledger statement with Running Balance."""
    cols = ["Txn ID", "Date", "Bill No", "Party / Company", "Product / Item",
            "Bill Amount", "Paid Amount", "Return Amount", "Running Balance", "Status"]
    sub = df if not party or party == "ALL PARTIES" else df[df["Party / Company"] == party]
    if sub.empty:
        return pd.DataFrame(columns=cols)
    sub = sub.sort_values(["Date", "Txn ID"]).copy()
    sub["Net Impact"] = sub["Bill Amount"] - sub["Paid Amount"] - sub["Return Amount"]
    sub["Running Balance"] = sub["Net Impact"].cumsum().round(2)
    return sub[cols].reset_index(drop=True)


# =====================================================================
# EXPORT (Excel + PDF)
# =====================================================================
def to_excel_bytes(sheets: dict[str, pd.DataFrame]) -> bytes:
    out = io.BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl", datetime_format="YYYY-MM-DD") as w:
        has_sheets = False
        for name, d in sheets.items():
            if not isinstance(d, pd.DataFrame):
                continue
            # Excel forbids  [ ] : * ? / \  in sheet names and limits them to 31 chars
            nm = (re.sub(r"[\[\]:*?/\\]", "-", str(name or "Sheet1")).strip() or "Sheet1")[:31]
            export_df = d.copy()
            # Normalize datetime columns to timezone-unaware formatted strings
            for col in export_df.columns:
                if pd.api.types.is_datetime64_any_dtype(export_df[col]):
                    export_df[col] = export_df[col].dt.strftime("%Y-%m-%d").fillna("")
            export_df.to_excel(w, sheet_name=nm, index=False)
            has_sheets = True
            ws = w.sheets[nm]
            if ws.max_column > 0:
                for col in ws.columns:
                    n = max((len(str(c.value)) if c.value is not None else 0) for c in col)
                    ws.column_dimensions[col[0].column_letter].width = min(max(n + 3, 10), 45)
                for cell in ws[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="1E3A8A")
                ws.freeze_panes = "A2"
        if not has_sheets:
            pd.DataFrame({"Message": ["No data"]}).to_excel(w, sheet_name="Data", index=False)
    return out.getvalue()


def _fmt(col: str, v) -> str:
    if v is None or (not isinstance(v, str) and pd.isna(v)):
        return ""
    if isinstance(v, (pd.Timestamp, datetime, date)):
        return v.strftime("%d %b %Y")
    if col in MONEY:
        try:
            return f"{float(v):,.0f}"
        except (ValueError, TypeError):
            return str(v)
    if col == "Qty":
        try:
            return f"{float(v):g}"
        except (ValueError, TypeError):
            return str(v)
    return str(v)


def to_pdf_bytes(title: str, df: pd.DataFrame) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=24, rightMargin=24,
                            topMargin=28, bottomMargin=28, title=title)
    ss = getSampleStyleSheet()
    fs = 8 if len(df.columns) <= 9 else 6.5
    cell = ParagraphStyle("c", parent=ss["Normal"], fontSize=fs, leading=fs + 2.5)
    cell_r = ParagraphStyle("cr", parent=cell, alignment=2)
    head = ParagraphStyle("h", parent=cell, textColor=colors.white, fontName="Helvetica-Bold")
    h1 = ParagraphStyle("t", parent=ss["Title"], fontSize=16, alignment=0, textColor=colors.HexColor("#1E3A8A"))
    sub = ParagraphStyle("s", parent=ss["Normal"], fontSize=9, textColor=colors.HexColor("#64748B"))

    if df.empty and len(df.columns) == 0:
        doc.build([Paragraph(escape(BUSINESS_NAME), h1),
                   Paragraph(f"{escape(title)} &nbsp;|&nbsp; Generated {datetime.now():%d %b %Y, %H:%M} &nbsp;|&nbsp; 0 record(s)", sub),
                   Spacer(1, 10), Paragraph("No records found.", ss["Normal"])])
        return buf.getvalue()

    right = [c in MONEY or c == "Qty" for c in df.columns]
    data = [[Paragraph(escape(str(c)), head) for c in df.columns]]
    for _, r in df.iterrows():
        data.append([Paragraph(escape(_fmt(c, v)), cell_r if right[i] else cell)
                     for i, (c, v) in enumerate(r.items())])
    lens = [min(max(max([len(str(c))] + [len(_fmt(c, v)) for v in df[c].head(200)]), 6), 28) for c in df.columns]
    tot_len = sum(lens) if sum(lens) > 0 else 1
    col_w = [doc.width * n / tot_len for n in lens] if lens else None
    t = Table(data, colWidths=col_w, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    doc.build([Paragraph(escape(BUSINESS_NAME), h1),
               Paragraph(f"{escape(title)} &nbsp;|&nbsp; Generated {datetime.now():%d %b %Y, %H:%M} "
                         f"&nbsp;|&nbsp; {len(df)} record(s)", sub),
               Spacer(1, 10), t])
    return buf.getvalue()


def invoice_pdf(row: pd.Series, items: pd.DataFrame) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=40, rightMargin=40, topMargin=40, bottomMargin=40,
                            title=f"Invoice {row['Bill No']}")
    ss = getSampleStyleSheet()
    navy = colors.HexColor("#1E3A8A")
    h1 = ParagraphStyle("a", parent=ss["Title"], fontSize=20, alignment=0, textColor=navy)
    norm = ParagraphStyle("n", parent=ss["Normal"], fontSize=10, leading=14)
    d = row["Date"].strftime("%d %b %Y") if pd.notna(row["Date"]) else "-"
    info = Table([[Paragraph(f"<b>Bill To</b><br/>{escape(str(row['Party / Company']))}", norm),
                   Paragraph(f"<b>Invoice No:</b> {escape(str(row['Bill No']))}<br/><b>Date:</b> {d}<br/>"
                             f"<b>Status:</b> {str(row['Status'])}", norm)]], colWidths=[280, 235])
    lines = [["Item", "Qty", "Rate", "Amount"]]
    if items.empty:
        lines.append([str(row["Product / Item"] or "Goods supplied"), f"{float(row['Qty']):g}" if row["Qty"] else "", "",
                      f"{float(row['Subtotal']):,.2f}"])
    for _, it in items.iterrows():
        lines.append([str(it["Product"]), f"{float(it['Qty']):g}", f"{float(it['Rate']):,.2f}", f"{float(it['Amount']):,.2f}"])
    t = Table(lines, colWidths=[255, 60, 90, 110], repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), navy), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    tot = [["Subtotal", f"{float(row['Subtotal']):,.2f}"]]
    if row.get("Discount", 0):
        tot.append(["Discount", f"- {float(row['Discount']):,.2f}"])
    if row.get("Tax", 0):
        tot.append(["Tax", f"{float(row['Tax']):,.2f}"])
    tot += [["Grand Total", f"{float(row['Bill Amount']):,.2f}"], ["Paid", f"{float(row['Paid Amount']):,.2f}"]]
    if row.get("Return Amount", 0):
        tot.append(["Returns", f"{float(row['Return Amount']):,.2f}"])
    tot.append(["Balance Due", f"{float(row['Balance']):,.2f}"])
    
    gt_idx = next((i for i, r in enumerate(tot) if r[0] == "Grand Total"), -1)
    tt = Table(tot, colWidths=[110, 110], hAlign="RIGHT")
    t_styles = [
        ("ALIGN", (1, 0), (1, -1), "RIGHT"),
        ("LINEABOVE", (0, -1), (-1, -1), 1, navy),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
    ]
    if gt_idx >= 0:
        t_styles.append(("FONTNAME", (0, gt_idx), (-1, gt_idx), "Helvetica-Bold"))
    tt.setStyle(TableStyle(t_styles))
    
    doc.build([Paragraph(escape(BUSINESS_NAME), h1), Paragraph("INVOICE", norm), Spacer(1, 14), info,
               Spacer(1, 16), t, Spacer(1, 12), tt, Spacer(1, 24),
               Paragraph(escape(str(row["Remarks"])) if row.get("Remarks") else "", norm)])
    return buf.getvalue()



def _norm(s) -> str:
    return re.sub(r"[^a-z0-9]", "", str(s).lower())


TX_ALIASES = {
    "txnid": "Txn ID", "date": "Date", "billdate": "Date", "transactiondate": "Date", "invoicedate": "Date",
    "billno": "Bill No", "billnumber": "Bill No", "invoice": "Bill No", "invoiceno": "Bill No",
    "party": "Party / Company", "partyname": "Party / Company", "partycompany": "Party / Company",
    "company": "Party / Company", "customer": "Party / Company", "customername": "Party / Company",
    "client": "Party / Company", "product": "Product / Item", "item": "Product / Item",
    "productitem": "Product / Item", "description": "Product / Item", "qty": "Qty", "quantity": "Qty",
    "rate": "Rate", "price": "Rate", "unitprice": "Rate", "subtotal": "Subtotal", "discount": "Discount",
    "tax": "Tax", "gst": "Tax", "amount": "Bill Amount", "billamount": "Bill Amount", "total": "Bill Amount",
    "totalamount": "Bill Amount", "paid": "Paid Amount", "paidamount": "Paid Amount", "received": "Paid Amount",
    "return": "Return Amount", "returns": "Return Amount", "returnamount": "Return Amount",
    "paymentmethod": "Payment Method", "method": "Payment Method", "mode": "Payment Method",
    "chequeno": "Cheque No", "chequenumber": "Cheque No", "chequedate": "Cheque Date", "bank": "Bank Name",
    "bankname": "Bank Name", "chequestatus": "Cheque Status", "salebase": "Sale Base", "remarks": "Remarks",
    "remark": "Remarks", "notes": "Remarks", "note": "Remarks",
}
PRODUCT_ALIASES = {"product": "Product", "productname": "Product", "item": "Product", "itemname": "Product",
                   "name": "Product", "unit": "Unit", "uom": "Unit", "rate": "Rate", "price": "Rate",
                   "unitprice": "Rate", "saleprice": "Rate"}


def read_upload(f) -> pd.DataFrame:
    name = f.name.lower()
    if name.endswith(".pdf"):
        try:
            import pdfplumber
        except ImportError as exc:
            raise ValueError("PDF import needs the 'pdfplumber' package (see requirements.txt).") from exc
        rows = []
        with pdfplumber.open(f) as pdf:
            for page in pdf.pages:
                for tbl in page.extract_tables():
                    rows += tbl
        if len(rows) < 2:
            raise ValueError("No data table found in this PDF. Scanned/photo PDFs cannot be read - "
                             "please use an Excel or CSV file instead.")
        head = [str(h or "").strip() or f"Column {i + 1}" for i, h in enumerate(rows[0])]
        body = [r for r in rows[1:] if [str(x or "").strip() for x in r] != [str(x or "").strip() for x in rows[0]]]
        return pd.DataFrame(body, columns=head)
    if name.endswith(".csv"):
        return pd.read_csv(f)
    xl = pd.ExcelFile(f)
    sheet = "Transactions" if "Transactions" in xl.sheet_names else xl.sheet_names[0]
    return xl.parse(sheet)


def prepare_import(raw: pd.DataFrame):
    ren = {c: TX_ALIASES[_norm(c)] for c in raw.columns if _norm(c) in TX_ALIASES}
    ignored = [str(c) for c in raw.columns if c not in ren]
    d = raw[list(ren)].rename(columns=ren)
    d = d.loc[:, ~d.columns.duplicated()]
    if "Party / Company" not in d.columns:
        raise ValueError("No Party / Company column found. Rename your party column to 'Party' and try again.")
    for c in ("Qty", "Rate", "Subtotal", "Discount", "Tax", "Bill Amount", "Paid Amount", "Return Amount"):
        d[c] = _num(d[c]) if c in d else 0.0
    for c in DATE_COLS:
        d[c] = _date(d[c]) if c in d else pd.NaT
    for c in TEXT_COLS:
        d[c] = _txt(d[c]) if c in d else ""
    d["Party / Company"] = d["Party / Company"].str.upper()
    d = d[d["Party / Company"] != ""].reset_index(drop=True)
    d["Bill Amount"] = d["Bill Amount"].where(d["Bill Amount"] > 0, d["Qty"] * d["Rate"])
    d["Payment Method"] = d["Payment Method"].replace("", "Not specified")
    return d, [c for c in dict.fromkeys(ren.values())], ignored



def _dup_key(f: pd.DataFrame) -> pd.Series:
    return (f["Party / Company"] + "|" + f["Bill No"] + "|" + f["Date"].dt.strftime("%Y-%m-%d").fillna("")
            + "|" + f["Bill Amount"].astype(float).round(2).map("{:.2f}".format))


def split_duplicates(d: pd.DataFrame, existing_tx: pd.DataFrame):
    """Drop rows that already exist (same party + bill no + date + amount)."""
    if existing_tx.empty:
        return d, 0
    dup = d["Bill No"].ne("") & _dup_key(d).isin(set(_dup_key(existing_tx)))
    return d[~dup].reset_index(drop=True), int(dup.sum())


def next_txn_ids(tx: pd.DataFrame, n: int) -> list[str]:
    nums = pd.to_numeric(tx["Txn ID"].astype(str).str.extract(r"(\d+)")[0], errors="coerce")
    start = int(nums.max()) if nums.notna().any() else 0
    return [f"TXN-{start + i + 1:06d}" for i in range(n)]


def next_bill_no(tx: pd.DataFrame) -> str:
    used, n = set(tx["Bill No"]), len(tx) + 1
    while f"B-{n:04d}" in used:
        n += 1
    return f"B-{n:04d}"


def commit_import(d: pd.DataFrame, add_products: bool, tx, items, products):
    """Merge imported rows into the frames. Returns (tx, items, products, rows_added)."""
    d = d.copy()
    d["Txn ID"] = next_txn_ids(tx, len(d))
    used = set(tx["Bill No"])
    n = len(tx) + 1
    for i in d.index[d["Bill No"] == ""]:
        while f"B-{n:04d}" in used:
            n += 1
        d.at[i, "Bill No"] = f"B-{n:04d}"
        used.add(d.at[i, "Bill No"])
    li = d[(d["Product / Item"] != "") & (d["Qty"] > 0) & (d["Rate"] > 0)]
    new_items = pd.DataFrame({"Txn ID": li["Txn ID"], "Product": li["Product / Item"], "Qty": li["Qty"],
                              "Rate": li["Rate"], "Amount": (li["Qty"] * li["Rate"]).round(2)})
    tx = recompute(pd.concat([tx, d], ignore_index=True))
    items = pd.concat([items, new_items], ignore_index=True)
    if add_products and not new_items.empty:
        p = new_items.drop_duplicates("Product", keep="last")
        p = pd.DataFrame({"Product": p["Product"], "Unit": "Pcs", "Rate": p["Rate"]})
        p = p[~p["Product"].str.lower().isin(products["Product"].astype(str).str.lower())]
        products = pd.concat([products, p], ignore_index=True)
    return tx, items, products, len(d)


def import_products(raw: pd.DataFrame, products: pd.DataFrame):
    """Merge a product list into products. Returns (products, rows_in_file)."""
    ren = {c: PRODUCT_ALIASES[_norm(c)] for c in raw.columns if _norm(c) in PRODUCT_ALIASES}
    d = raw[list(ren)].rename(columns=ren)
    d = d.loc[:, ~d.columns.duplicated()]
    if "Product" not in d.columns:
        raise ValueError("No Product / Item column found in this file.")
    d["Product"] = _txt(d["Product"])
    d["Unit"] = _txt(d["Unit"]).replace("", "Pcs") if "Unit" in d else "Pcs"
    d["Rate"] = _num(d["Rate"]) if "Rate" in d else 0.0
    d = d[d["Product"] != ""][PRODUCT_COLS]
    products = pd.concat([products, d]).drop_duplicates("Product", keep="last").reset_index(drop=True)
    return products, len(d)


def ensure_txn_ids(tx: pd.DataFrame) -> pd.DataFrame:
    """Make every Txn ID filled and unique (needed for the database primary key)."""
    tx = tx.copy()
    bad = tx["Txn ID"].eq("") | tx["Txn ID"].duplicated(keep="first")
    if bad.any():
        nums = pd.to_numeric(tx.loc[~bad, "Txn ID"].str.extract(r"(\d+)")[0], errors="coerce")
        start = int(nums.max()) if nums.notna().any() else 0
        tx.loc[bad, "Txn ID"] = [f"TXN-{start + i + 1:06d}" for i in range(int(bad.sum()))]
    return tx


import urllib.parse


def whatsapp_invoice_text(row: pd.Series) -> str:
    party = str(row.get("Party / Company", "Valued Customer"))
    bill_no = str(row.get("Bill No", "-"))
    d_val = row.get("Date")
    d_str = d_val.strftime("%d %b %Y") if pd.notna(d_val) and hasattr(d_val, "strftime") else str(d_val or "-")
    total = float(row.get("Bill Amount", 0))
    paid = float(row.get("Paid Amount", 0))
    balance = float(row.get("Balance", 0))
    status = str(row.get("Status", "UNPAID"))

    msg = (
        f"Assalam-o-Alaikum / Dear {party},\n\n"
        f"This is an update regarding Invoice #{bill_no} from {BUSINESS_NAME}.\n\n"
        f"📅 Date: {d_str}\n"
        f"💰 Invoice Total: {CUR} {total:,.0f}\n"
        f"✅ Paid Amount: {CUR} {paid:,.0f}\n"
        f"⚠️ Balance Due: {CUR} {balance:,.0f}\n"
        f"📌 Payment Status: {status}\n\n"
        f"Thank you for your valuable business!\n"
        f"Regards,\n{BUSINESS_NAME}"
    )
    return msg


def whatsapp_url(phone: str, text: str) -> str:
    cleaned = re.sub(r"[^\d]", "", str(phone or ""))
    encoded = urllib.parse.quote(text)
    if cleaned:
        return f"https://wa.me/{cleaned}?text={encoded}"
    return f"https://wa.me/?text={encoded}"


def thermal_receipt_html(row: pd.Series, items: pd.DataFrame) -> str:
    d_val = row.get("Date")
    d_str = d_val.strftime("%d %b %Y") if pd.notna(d_val) and hasattr(d_val, "strftime") else str(d_val or "-")
    party = escape(str(row.get("Party / Company", "")))
    bill_no = escape(str(row.get("Bill No", "")))
    status = escape(str(row.get("Status", "")))
    subtotal = float(row.get("Subtotal", 0))
    discount = float(row.get("Discount", 0))
    tax = float(row.get("Tax", 0))
    total = float(row.get("Bill Amount", 0))
    paid = float(row.get("Paid Amount", 0))
    returns = float(row.get("Return Amount", 0))
    balance = float(row.get("Balance", 0))

    items_rows = ""
    if not items.empty:
        for _, it in items.iterrows():
            items_rows += f"""
            <tr>
                <td style="text-align:left;padding:4px 0;">{escape(str(it.get('Product', '')))}</td>
                <td style="text-align:center;padding:4px 0;">{float(it.get('Qty', 0)):g}</td>
                <td style="text-align:right;padding:4px 0;">{float(it.get('Rate', 0)):,.0f}</td>
                <td style="text-align:right;padding:4px 0;font-weight:600;">{float(it.get('Amount', 0)):,.0f}</td>
            </tr>
            """
    else:
        prod = escape(str(row.get("Product / Item", "Goods/Services")))
        qty = float(row.get("Qty", 1))
        items_rows = f"""
        <tr>
            <td style="text-align:left;padding:4px 0;">{prod}</td>
            <td style="text-align:center;padding:4px 0;">{qty:g}</td>
            <td style="text-align:right;padding:4px 0;">-</td>
            <td style="text-align:right;padding:4px 0;font-weight:600;">{subtotal:,.0f}</td>
        </tr>
        """

    disc_row = f"<div style='display:flex;justify-content:space-between;padding:2px 0;'><span>Discount:</span><span>- {CUR} {discount:,.0f}</span></div>" if discount > 0 else ""
    tax_row = f"<div style='display:flex;justify-content:space-between;padding:2px 0;'><span>Tax:</span><span>+ {CUR} {tax:,.0f}</span></div>" if tax > 0 else ""
    ret_row = f"<div style='display:flex;justify-content:space-between;padding:2px 0;'><span>Returned:</span><span>{CUR} {returns:,.0f}</span></div>" if returns > 0 else ""

    html = f"""
    <div style="max-width:320px;margin:12px auto;padding:16px;background:#ffffff;color:#1e293b;border:1px dashed #cbd5e1;border-radius:8px;font-family:'Courier New', Courier, monospace;font-size:12px;line-height:1.4;">
        <div style="text-align:center;margin-bottom:10px;">
            <div style="font-size:15px;font-weight:800;letter-spacing:1px;text-transform:uppercase;">{escape(BUSINESS_NAME)}</div>
            <div style="font-size:11px;color:#64748b;">SALE RECEIPT / INVOICE</div>
        </div>
        <div style="border-top:1px dashed #94a3b8;border-bottom:1px dashed #94a3b8;padding:6px 0;margin:8px 0;font-size:11px;">
            <div><b>Bill No:</b> {bill_no}</div>
            <div><b>Date:</b> {d_str}</div>
            <div><b>Party:</b> {party}</div>
            <div><b>Status:</b> {status}</div>
        </div>
        <table style="width:100%;font-size:11px;border-collapse:collapse;margin:8px 0;">
            <thead>
                <tr style="border-bottom:1px solid #cbd5e1;text-transform:uppercase;color:#475569;">
                    <th style="text-align:left;padding:4px 0;">Item</th>
                    <th style="text-align:center;padding:4px 0;">Qty</th>
                    <th style="text-align:right;padding:4px 0;">Rate</th>
                    <th style="text-align:right;padding:4px 0;">Amt</th>
                </tr>
            </thead>
            <tbody>
                {items_rows}
            </tbody>
        </table>
        <div style="border-top:1px dashed #94a3b8;padding-top:6px;font-size:11px;">
            <div style="display:flex;justify-content:space-between;padding:2px 0;"><span>Subtotal:</span><span>{CUR} {subtotal:,.0f}</span></div>
            {disc_row}
            {tax_row}
            <div style="display:flex;justify-content:space-between;padding:4px 0;font-size:13px;font-weight:800;border-top:1px solid #cbd5e1;margin-top:4px;">
                <span>Total:</span><span>{CUR} {total:,.0f}</span>
            </div>
            <div style="display:flex;justify-content:space-between;padding:2px 0;"><span>Paid:</span><span>{CUR} {paid:,.0f}</span></div>
            {ret_row}
            <div style="display:flex;justify-content:space-between;padding:4px 0;font-size:13px;font-weight:800;color:#0f172a;background:#f1f5f9;border-radius:4px;padding:4px 6px;margin-top:4px;">
                <span>Balance Due:</span><span>{CUR} {balance:,.0f}</span>
            </div>
        </div>
        <div style="text-align:center;margin-top:14px;padding-top:8px;border-top:1px dashed #cbd5e1;font-size:10px;color:#64748b;">
            Thank you for your business!<br/>Computer generated receipt
        </div>
    </div>
    """
    return html
